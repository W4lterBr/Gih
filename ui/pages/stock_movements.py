# -*- coding: utf-8 -*-
"""Página de Movimentação de estoque (entradas, saídas, perdas)."""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any, Callable, Optional, cast

from PyQt6.QtGui import QColor, QFont
from PyQt6.QtWidgets import (
    QAbstractItemView, QComboBox, QDialog, QFileDialog, QHBoxLayout,
    QHeaderView, QLabel, QLineEdit, QPushButton, QTableWidget, QTableWidgetItem,
    QVBoxLayout,
)

from core.database import Database
from core.formatting import format_datetime, money
from core.stock_categories import MOVEMENT_CATEGORY_COLORS
from ui.dialogs.stock_movement_dialog import StockMovementDialog
from ui.messages import show_message
from ui.pages.base import BasePage
from ui.widgets.stat_card import StatCard


class StockMovementsPage(BasePage):
    """Movimentação de estoque: histórico completo de entradas, saídas e produção,
    com registro manual (inclusive perdas/quebras - antes impossível de registrar)."""

    FILTERS: list[tuple[str, str]] = [
        ("todas", "📋 Todas"),
        ("entrada", "📥 Entradas"),
        ("saida", "📤 Saídas"),
        ("producao", "🏭 Produção"),
        ("perda", "💔 Perdas"),
    ]

    def __init__(self, db: Database, toast_cb: Optional[Callable[[str], None]] = None) -> None:
        super().__init__("Movimentação", "Entradas, saídas, produção e perdas de estoque")
        self.db = db
        self.toast_cb = toast_cb
        self.current_filter = "todas"
        bl = QVBoxLayout(self.body)

        # --- Cards de resumo ---
        self.cards_box = QHBoxLayout()
        self.card_entradas = StatCard("Entradas hoje", "0", "#10B981", "ph.arrow-circle-down")
        self.card_saidas = StatCard("Saídas hoje", "0", "#3B82F6", "ph.arrow-circle-up")
        self.card_producao = StatCard("Produção (mês)", "0", "#8B5CF6", "ph.cooking-pot")
        self.card_perdas = StatCard("Perdas (mês)", "0 (R$ 0,00)", "#EF4444", "ph.warning")
        for card in (self.card_entradas, self.card_saidas, self.card_producao, self.card_perdas):
            self.cards_box.addWidget(card)
        bl.addLayout(self.cards_box)

        # --- Ações ---
        actions = QHBoxLayout()
        self.btn_add = QPushButton("+ Nova Movimentação")
        self.btn_add.setProperty("primary", True)
        cast(Any, self.btn_add.clicked).connect(self._add_movement)
        actions.addWidget(self.btn_add)
        self.btn_del = QPushButton("🗑️ Excluir")
        cast(Any, self.btn_del.clicked).connect(self._delete_movement)
        actions.addWidget(self.btn_del)
        self.btn_export = QPushButton("📊 Exportar Excel")
        cast(Any, self.btn_export.clicked).connect(self._export_excel)
        actions.addWidget(self.btn_export)
        actions.addStretch(1)
        bl.addLayout(actions)

        # --- Filtros (chips) + busca + período ---
        from PyQt6.QtWidgets import QButtonGroup
        filters_box = QHBoxLayout()
        self.filter_group = QButtonGroup(self)
        self.filter_group.setExclusive(True)
        self.filter_buttons: dict[str, QPushButton] = {}
        for key, label_txt in self.FILTERS:
            btn = QPushButton(label_txt)
            btn.setCheckable(True)
            btn.setProperty("chip", True)
            self.filter_group.addButton(btn)
            self.filter_buttons[key] = btn
            filters_box.addWidget(btn)
        self.filter_buttons["todas"].setChecked(True)
        cast(Any, self.filter_group.buttonClicked).connect(self._on_filter_clicked)
        filters_box.addStretch(1)

        filters_box.addWidget(QLabel("Período:"))
        self.period_combo = QComboBox()
        self.period_combo.addItems(["Hoje", "7 dias", "30 dias", "Tudo"])
        self.period_combo.setCurrentText("30 dias")
        cast(Any, self.period_combo.currentTextChanged).connect(lambda _t: self.refresh())
        filters_box.addWidget(self.period_combo)
        bl.addLayout(filters_box)

        search_box = QHBoxLayout()
        search_box.addWidget(QLabel("Pesquisar produto:"))
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Nome do produto…")
        self.search_edit.setClearButtonEnabled(True)
        cast(Any, self.search_edit.textChanged).connect(lambda _t: self.refresh())
        search_box.addWidget(self.search_edit, 1)
        bl.addLayout(search_box)

        # --- Tabela ---
        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels(["Data", "Produto", "Tipo", "Categoria", "Quantidade", "Motivo", "Usuário"])
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        if header := self.table.horizontalHeader():
            header.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
            header.setStretchLastSection(True)
        bl.addWidget(self.table)

        self.refresh()

        # Atualiza a cada 5s enquanto visível (estoque pode mudar pelo painel da cozinha)
        self.start_auto_refresh(self.refresh, 5000)

    def _on_filter_clicked(self, _btn: QPushButton) -> None:
        for key, btn in self.filter_buttons.items():
            if btn.isChecked():
                self.current_filter = key
                break
        self.refresh()

    def _period_clause(self) -> str:
        period = self.period_combo.currentText()
        if period == "Hoje":
            return "DATE(sm.created_at) = DATE('now','localtime')"
        if period == "7 dias":
            return "DATE(sm.created_at) >= DATE('now','localtime','-7 days')"
        if period == "30 dias":
            return "DATE(sm.created_at) >= DATE('now','localtime','-30 days')"
        return "1=1"

    def _filter_clause(self) -> str:
        if self.current_filter == "entrada":
            return "sm.type = 'entrada'"
        if self.current_filter == "saida":
            return "sm.type = 'saida'"
        if self.current_filter == "producao":
            return "(sm.category = 'Produção' OR (sm.category IS NULL AND sm.reason LIKE 'Produção%'))"
        if self.current_filter == "perda":
            return "sm.category = 'Perda/Quebra'"
        return "1=1"

    def _refresh_summary_cards(self) -> None:
        try:
            row = self.db.query("""
                SELECT
                    COALESCE(SUM(CASE WHEN type='entrada' AND DATE(created_at)=DATE('now','localtime') THEN quantity ELSE 0 END), 0) AS entradas_hoje,
                    COALESCE(SUM(CASE WHEN type='saida' AND DATE(created_at)=DATE('now','localtime') THEN quantity ELSE 0 END), 0) AS saidas_hoje,
                    COALESCE(SUM(CASE WHEN (category='Produção' OR (category IS NULL AND reason LIKE 'Produção%'))
                                    AND strftime('%Y-%m', created_at)=strftime('%Y-%m','now','localtime') THEN quantity ELSE 0 END), 0) AS producao_mes
                FROM stock_movements
            """)[0]
            self.card_entradas.value_label.setText(str(int(row["entradas_hoje"])))  # type: ignore[attr-defined]
            self.card_saidas.value_label.setText(str(int(row["saidas_hoje"])))  # type: ignore[attr-defined]
            self.card_producao.value_label.setText(str(int(row["producao_mes"])))  # type: ignore[attr-defined]

            perda_row = self.db.query("""
                SELECT
                    COALESCE(SUM(sm.quantity), 0) AS qtd,
                    COALESCE(SUM(sm.quantity * COALESCE(p.price, 0)), 0) AS valor
                FROM stock_movements sm
                LEFT JOIN products p ON p.id = sm.product_id
                WHERE sm.category = 'Perda/Quebra'
                  AND strftime('%Y-%m', sm.created_at) = strftime('%Y-%m','now','localtime')
            """)[0]
            self.card_perdas.value_label.setText(  # type: ignore[attr-defined]
                f"{int(perda_row['qtd'])} ({money(float(perda_row['valor']))})"
            )
        except Exception as e:
            print(f"Erro ao atualizar cards de movimentação: {e}")

    def refresh(self) -> None:
        try:
            self._refresh_summary_cards()

            term = self.search_edit.text().strip()
            where_parts = [self._period_clause(), self._filter_clause()]
            params: list[Any] = []
            if term:
                where_parts.append("p.name LIKE ?")
                params.append(f"%{term}%")
            where_sql = " AND ".join(where_parts)

            rows = self.db.query(f"""
                SELECT sm.id, sm.type, sm.category, sm.quantity, sm.reason,
                       sm.created_at, sm.created_by, sm.order_id, p.name AS product_name
                FROM stock_movements sm
                JOIN products p ON p.id = sm.product_id
                WHERE {where_sql}
                ORDER BY sm.created_at DESC, sm.id DESC
                LIMIT 500
            """, tuple(params))

            if not self.rows_changed("table", rows):
                return  # nada mudou: mantém a tabela (rolagem e seleção)
            self.table.setRowCount(0)
            for r in rows:
                row = self.table.rowCount()
                self.table.insertRow(row)
                self.table.setItem(row, 0, QTableWidgetItem(format_datetime(r["created_at"])))
                self.table.setItem(row, 1, QTableWidgetItem(r["product_name"]))

                tipo_txt = "Entrada" if r["type"] == "entrada" else "Saída"
                tipo_item = QTableWidgetItem(tipo_txt)
                tipo_item.setForeground(QColor("#10B981" if r["type"] == "entrada" else "#3B82F6"))
                tipo_item.setFont(QFont("Arial", 9, QFont.Weight.Bold))
                self.table.setItem(row, 2, tipo_item)

                categoria_txt = r["category"] or "—"
                cat_item = QTableWidgetItem(categoria_txt)
                cat_color = MOVEMENT_CATEGORY_COLORS.get(categoria_txt)
                if cat_color:
                    cat_item.setForeground(QColor(cat_color))
                self.table.setItem(row, 3, cat_item)

                self.table.setItem(row, 4, QTableWidgetItem(str(r["quantity"])))
                self.table.setItem(row, 5, QTableWidgetItem(r["reason"] or ""))
                self.table.setItem(row, 6, QTableWidgetItem(r["created_by"] or "—"))
                self.table.setVerticalHeaderItem(row, QTableWidgetItem(str(r["id"])))
        except Exception as e:
            print(f"Erro no refresh de Movimentação: {e}")

    def _add_movement(self) -> None:
        dlg = StockMovementDialog(self.db, self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        try:
            values = dlg.get_values()
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            # Usa a data escolhida no diálogo mas preserva a hora atual, para manter
            # a ordenação por created_at coerente mesmo em lançamentos retroativos.
            movement_date = f"{values['date']} {now.split(' ')[1]}"
            user = getattr(self.db, "current_user", None) or "unknown"

            # Ajusta o estoque ANTES de gravar o movimento, e para saída usa um UPDATE
            # condicional (WHERE stock >= ?) checando quantas linhas foram afetadas -
            # a validação do diálogo lê o estoque só quando ele abriu, então duas
            # movimentações quase simultâneas (ou o refresh automático da tela) podiam
            # passar na validação com o mesmo valor "antigo" e ambas serem aplicadas,
            # derrubando o estoque abaixo de zero. Esse UPDATE condicional fecha essa
            # janela: só uma das duas consegue de fato decrementar.
            if values["type"] == "entrada":
                self.db.execute("UPDATE products SET stock = stock + ? WHERE id=?",
                                 (values["quantity"], values["product_id"]))
            else:
                cur_stock = self.db.execute(
                    "UPDATE products SET stock = stock - ? WHERE id=? AND stock >= ?",
                    (values["quantity"], values["product_id"], values["quantity"])
                )
                if cur_stock.rowcount == 0:
                    current = self.db.query("SELECT stock FROM products WHERE id=?", (values["product_id"],))
                    available = current[0]["stock"] if current else 0
                    show_message(
                        self, "Estoque insuficiente",
                        f"Estoque disponível agora: {available} un. Não é possível dar saída de "
                        f"{values['quantity']} un - outra movimentação pode ter alterado o estoque "
                        f"enquanto este formulário estava aberto.",
                        ("OK",)
                    )
                    return

            cur = self.db.execute(
                """INSERT INTO stock_movements(product_id, type, category, quantity, reason, order_id, created_at, created_by)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (values["product_id"], values["type"], values["category"], values["quantity"],
                 values["notes"], None, movement_date, str(user))
            )
            movement_id = cur.lastrowid

            try:
                self.db.audit_log(
                    "stock_movement", movement_id, "create",
                    details=f"{values['type']} | {values['category']} | produto={values['product_name']} | qtd={values['quantity']}"
                )
            except Exception:
                pass

            self.refresh()
            if self.toast_cb:
                verbo = "adicionadas" if values["type"] == "entrada" else "removidas"
                self.toast_cb(f"✅ {values['quantity']} un {verbo} ({values['category']}) - {values['product_name']}")
        except Exception as e:
            show_message(self, "Erro ao registrar movimentação", str(e), ("OK",))

    def _selected_movement_ids(self) -> list[int]:
        # Antes só existia _selected_movement_id() (currentRow(), uma linha), então
        # selecionar várias movimentações e clicar Excluir só apagava a última focada.
        ids: list[int] = []
        for index in self.table.selectionModel().selectedRows():
            vh = self.table.verticalHeaderItem(index.row())
            if vh:
                try:
                    ids.append(int(vh.text()))
                except (ValueError, AttributeError):
                    pass
        return list(dict.fromkeys(ids))

    def _delete_movement(self) -> None:
        mov_ids = self._selected_movement_ids()
        if not mov_ids:
            if self.toast_cb:
                self.toast_cb("Selecione uma ou mais movimentações para excluir")
            return
        try:
            count = len(mov_ids)
            msg = "Excluir esta movimentação? O estoque será ajustado de volta." if count == 1 \
                else f"Excluir {count} movimentações selecionadas? O estoque de cada produto será ajustado de volta."
            res = show_message(self, "Confirmação", msg, ("Não", "Sim"), default=1)
            if res != 1:
                return

            deleted = 0
            blocked = 0
            # Agrupa tudo em uma única transação para não travar a janela numa seleção grande
            with self.db.transaction():
                for mov_id in mov_ids:
                    row = self.db.query("SELECT product_id, type, quantity, order_id, category FROM stock_movements WHERE id=?", (mov_id,))
                    if not row:
                        continue
                    data = row[0]
                    if data["order_id"] is not None:
                        blocked += 1
                        continue

                    # Reverte o efeito no estoque (simétrico ao que foi feito na criação)
                    reverse_delta = -data["quantity"] if data["type"] == "entrada" else data["quantity"]
                    self.db.execute("UPDATE products SET stock = stock + ? WHERE id=?", (reverse_delta, data["product_id"]))
                    self.db.execute("DELETE FROM stock_movements WHERE id=?", (mov_id,))
                    deleted += 1
                    try:
                        self.db.audit_log("stock_movement", mov_id, "delete", details=f"revertido: {data['type']} {data['quantity']}un")
                    except Exception:
                        pass

            self.refresh()
            if self.toast_cb:
                if deleted and not blocked:
                    self.toast_cb("🗑️ Movimentação excluída e estoque ajustado" if deleted == 1
                                   else f"🗑️ {deleted} movimentações excluídas e estoque ajustado")
                elif deleted and blocked:
                    self.toast_cb(f"🗑️ {deleted} excluída(s); {blocked} bloqueada(s) por serem geradas por pedido")
                elif blocked:
                    show_message(
                        self, "Não é possível excluir",
                        "A(s) movimentação(ões) selecionada(s) foi(ram) gerada(s) automaticamente por um pedido "
                        "e não pode(m) ser excluída(s) aqui - gerencie pelo pedido correspondente na tela de Pedidos.",
                        ("OK",)
                    )
        except Exception as e:
            show_message(self, "Erro ao excluir", str(e), ("OK",))

    def _export_excel(self) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            term = self.search_edit.text().strip()
            where_parts = [self._period_clause(), self._filter_clause()]
            params: list[Any] = []
            if term:
                where_parts.append("p.name LIKE ?")
                params.append(f"%{term}%")
            where_sql = " AND ".join(where_parts)

            rows = self.db.query(f"""
                SELECT sm.created_at, p.name AS product_name, sm.type, sm.category,
                       sm.quantity, sm.reason, sm.created_by
                FROM stock_movements sm
                JOIN products p ON p.id = sm.product_id
                WHERE {where_sql}
                ORDER BY sm.created_at DESC
            """, tuple(params))

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Movimentações"
            headers = ["Data", "Produto", "Tipo", "Categoria", "Quantidade", "Motivo", "Usuário"]
            ws.append(headers)
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for r in rows:
                ws.append([
                    format_datetime(r["created_at"]),
                    r["product_name"],
                    "Entrada" if r["type"] == "entrada" else "Saída",
                    r["category"] or "",
                    r["quantity"],
                    r["reason"] or "",
                    r["created_by"] or "",
                ])

            for col_cells in ws.columns:
                length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 12), 40)

            default_name = f"movimentacoes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            path, _ = QFileDialog.getSaveFileName(self, "Exportar Movimentações", default_name, "Excel (*.xlsx)")
            if not path:
                return
            wb.save(path)
            if self.toast_cb:
                self.toast_cb(f"📊 Exportado: {os.path.basename(path)}")
        except Exception as e:
            show_message(self, "Erro ao exportar", str(e), ("OK",))
