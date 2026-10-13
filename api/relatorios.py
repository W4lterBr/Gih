# -*- coding: utf-8 -*-
"""Relatórios: Entregas (dia), Produção (movimentações) e Fechamento (por cliente)."""

from __future__ import annotations

from collections import OrderedDict
from datetime import datetime

from api import ErroApi, Pedido, contexto, rota
from api.pedidos import _logo_empresa, separar_lote
from api.util import data_iso, hoje_iso, qtd_texto


def _entregas(dia: str) -> dict:
    rows = contexto.db().query("""
        SELECT c.name AS cliente, o.quantity AS quantidade, p.unit, p.name AS produto, c.address AS endereco,
               c.phone AS telefone, o.notes, o.status, COALESCE(o.label,'Comum') AS etiqueta
        FROM orders o JOIN customers c ON c.id = o.customer_id JOIN products p ON p.id = o.product_id
        WHERE DATE(o.delivery_date) = ? AND COALESCE(o.situacao, 'aberto') != 'cancelado'
        ORDER BY c.name
    """, (dia,))
    # Todos os pedidos do dia que não foram cancelados: "Pagamento na Retirada" e
    # "Aguardando" também são entregues (antes só os pagos entravam). O pagamento vai
    # numa coluna para o entregador saber o que cobrar.
    linhas = [{**dict(r), "obs": separar_lote(r["notes"])[2]} for r in rows]
    return {
        "titulo": "Relatório de Entregas",
        "meta": f"Dia: {datetime.fromisoformat(dia).strftime('%d/%m/%Y')} | Pedidos não cancelados",
        "colunas": ["Cliente", "Quantidade", "Produto", "Endereço", "Pagamento", "Observação"],
        "linhas": [[x["cliente"], qtd_texto(x["quantidade"], x["unit"]), x["produto"], x["endereco"] or "", x["status"], x["obs"] or ""] for x in linhas],
        "rodape": f"Total: {_total_misto(linhas, 'quantidade')} | Itens: {len(linhas)}",
    }


def _producao(de: str, ate: str) -> dict:
    rows = contexto.db().query("""
        SELECT sm.created_at AS data, p.name AS produto, p.unit, sm.type AS tipo, sm.category AS categoria,
               sm.quantity AS quantidade, sm.reason AS motivo
        FROM stock_movements sm JOIN products p ON p.id = sm.product_id
        WHERE DATE(sm.created_at) BETWEEN ? AND ?
        ORDER BY sm.created_at DESC
    """, (de, ate))
    entradas = _total_misto([r for r in rows if r["tipo"] == "entrada"], "quantidade")
    saidas = _total_misto([r for r in rows if r["tipo"] == "saida"], "quantidade")
    return {
        "titulo": "Relatório de Produção",
        "meta": f"Período: {_br(de)} a {_br(ate)}",
        "colunas": ["Data", "Produto", "Tipo", "Categoria", "Quantidade", "Motivo"],
        "linhas": [[_br(r["data"][:10]), r["produto"], "Entrada" if r["tipo"] == "entrada" else "Saída",
                    r["categoria"] or "", qtd_texto(r["quantidade"], r["unit"]), r["motivo"] or ""] for r in rows],
        "rodape": f"Entradas: {entradas} | Saídas: {saidas}",
    }


def _fechamento(de: str, ate: str, etiqueta: str) -> dict:
    sql = """
        SELECT c.name AS cliente, COALESCE(o.label,'Comum') AS etiqueta, p.name AS produto, p.unit,
               DATE(o.delivery_date) AS data, o.quantity AS quantidade, o.quantity * COALESCE(p.price,0) AS valor
        FROM orders o JOIN customers c ON c.id=o.customer_id JOIN products p ON p.id=o.product_id
        WHERE DATE(o.delivery_date) BETWEEN ? AND ? AND COALESCE(o.situacao, 'aberto') != 'cancelado'"""
    params: list = [de, ate]
    if etiqueta:
        sql += " AND LOWER(COALESCE(o.label,'Comum')) = LOWER(?)"
        params.append(etiqueta)
    rows = contexto.db().query(sql + " ORDER BY c.name, etiqueta, data", tuple(params))
    grupos: OrderedDict = OrderedDict()
    for r in rows:
        k = (r["cliente"], r["etiqueta"], r["produto"])
        g = grupos.setdefault(k, {"data": r["data"], "quantidade": 0, "valor": 0.0, "unit": r["unit"]})
        g["quantidade"] += r["quantidade"]
        g["valor"] += r["valor"] or 0
    linhas = [[k[0], k[1], k[2], _br(g["data"]), qtd_texto(round(g["quantidade"], 3), g["unit"]), _moeda(g["valor"])]
              for k, g in sorted(grupos.items())]
    resumo: OrderedDict = OrderedDict()
    for (cli, eti, _), g in sorted(grupos.items()):
        s = resumo.setdefault((cli, eti), [0, [], 0.0])
        s[0] += 1
        s[1].append({"quantidade": g["quantidade"], "unit": g["unit"]})
        s[2] += g["valor"]
    clientes = {k[0] for k in grupos}
    total_qtd = _total_misto(list(grupos.values()), "quantidade")
    total_valor = sum(g["valor"] for g in grupos.values())
    return {
        "titulo": "Relatório de Fechamento",
        "meta": f"Período: {_br(de)} a {_br(ate)} | Etiqueta: {etiqueta or 'Todas'} | Valor pelo preço de cadastro do produto",
        "colunas": ["Cliente", "Etiqueta", "Produto", "Data", "Quantidade", "Valor"],
        "linhas": linhas,
        "rodape": f"Clientes: {len(clientes)} | Itens: {len(linhas)} | Quantidade: {total_qtd} | Valor: {_moeda(total_valor)}",
        "resumo": {"colunas": ["Cliente", "Etiqueta", "Produtos", "Quantidade", "Valor"],
                   "linhas": [[c, e, s[0], _total_misto(s[1], "quantidade"), _moeda(s[2])] for (c, e), s in resumo.items()]},
    }


def _total_misto(linhas, campo: str) -> str:
    """Soma separada por unidade: "12 un + 3,5 kg" (somar unidades com quilos não faz sentido)."""
    un = sum(int(x[campo] or 0) for x in linhas if (x["unit"] or "un") != "kg")
    kg = round(sum(float(x[campo] or 0) for x in linhas if x["unit"] == "kg"), 3)
    partes = []
    if un or not kg:
        partes.append(f"{un} un")
    if kg:
        partes.append(qtd_texto(kg, "kg"))
    return " + ".join(partes)


def _br(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso[:10]).strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return iso or ""


def _moeda(v: float) -> str:
    from core.formatting import money
    return money(v or 0)


def _montar(p: Pedido, fonte: dict) -> dict:
    tipo = fonte.get("tipo")
    if tipo == "entregas":
        return _entregas(data_iso(fonte.get("dia") or hoje_iso(), "dia"))
    de = data_iso(fonte.get("de") or hoje_iso()[:8] + "01", "de")
    ate = data_iso(fonte.get("ate") or hoje_iso(), "ate")
    if tipo == "producao":
        return _producao(de, ate)
    if tipo == "fechamento":
        return _fechamento(de, ate, (fonte.get("etiqueta") or "").strip())
    raise ErroApi("Relatório desconhecido.")


@rota("GET", "/relatorios")
def relatorio(p: Pedido) -> dict:
    return _montar(p, p.query)


@rota("POST", "/relatorios/exportar", escreve=False)
def exportar(p: Pedido) -> dict:
    dados = _montar(p, p.corpo)
    caminho = str(p.c("caminho") or "")
    empresa, logo = _logo_empresa(contexto.db())
    if caminho.lower().endswith(".pdf"):
        _pdf(caminho, empresa, logo, dados)
    elif caminho.lower().endswith(".xlsx"):
        _xlsx(caminho, empresa, logo, dados)
    else:
        raise ErroApi("Escolha um arquivo .pdf ou .xlsx.")
    return {"caminho": caminho}


def _pdf(caminho, empresa, logo, d):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    doc = SimpleDocTemplate(caminho, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=40, bottomMargin=36, title=d["titulo"])
    est = getSampleStyleSheet()
    story = []
    if logo:
        try:
            from reportlab.lib.utils import ImageReader
            w, h = ImageReader(logo).getSize()
            esc = min(120 / w, 50 / h)
            img = Image(logo, width=w * esc, height=h * esc)
            img.hAlign = "LEFT"
            story.append(img)
        except Exception:
            pass
    story += [Paragraph(f"<b>{escape(empresa)}</b> · {escape(d['titulo'])}", est["Title"]),
              Paragraph(escape(d["meta"]), est["Italic"]), Spacer(1, 12)]

    def tabela(cab, linhas, cor):
        t = Table([cab] + [[Paragraph(escape(str(v)), est["BodyText"]) for v in l] for l in linhas], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(cor)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return t

    story.append(tabela(d["colunas"], d["linhas"], "#2563EB"))
    story += [Spacer(1, 8), Paragraph(escape(d["rodape"]), est["Normal"])]
    if d.get("resumo"):
        story += [Spacer(1, 18), Paragraph("<b>Resumo por Cliente/Etiqueta</b>", est["Heading3"]),
                  tabela(d["resumo"]["colunas"], d["resumo"]["linhas"], "#1F2937")]
    doc.build(story)


def _xlsx(caminho, empresa, logo, d):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = d["titulo"][:31]
    lin = 1
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(logo)
            esc = min(120 / img.width, 60 / img.height)
            img.width, img.height = img.width * esc, img.height * esc
            ws.add_image(img, "A1")
            lin = 5
        except Exception:
            pass
    ws.cell(lin, 1, f"{empresa} · {d['titulo']}").font = Font(bold=True, size=14)
    ws.cell(lin + 1, 1, d["meta"]).font = Font(italic=True, size=10)
    fina = Side(style="thin", color="D1D5DB")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)

    def bloco(inicio, cab, linhas, cor):
        for j, nome in enumerate(cab, 1):
            c = ws.cell(inicio, j, nome)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=cor)
            c.alignment, c.border = Alignment(horizontal="center"), borda
        for i, l in enumerate(linhas, 1):
            for j, v in enumerate(l, 1):
                ws.cell(inicio + i, j, v).border = borda
        return inicio + len(linhas) + 1

    fim = bloco(lin + 3, d["colunas"], d["linhas"], "2563EB")
    ws.cell(fim + 1, 1, d["rodape"]).font = Font(bold=True)
    if d.get("resumo"):
        ws.cell(fim + 3, 1, "Resumo por Cliente/Etiqueta").font = Font(bold=True, size=12)
        bloco(fim + 4, d["resumo"]["colunas"], d["resumo"]["linhas"], "1F2937")
    for col in ws.columns:
        largura = max((len(str(c.value)) for c in col if c.value is not None and c.row > lin + 1), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(largura + 2, 10), 50)
    wb.save(caminho)
