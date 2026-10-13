# -*- coding: utf-8 -*-
"""Pedidos: listagem (com agrupamento de lotes), criação, edição, exclusão e relatório.

Formato no banco igual ao da versão 1.x (bancos existentes continuam valendo) e ao que o
painel da cozinha lê:
- uma linha em `orders` por produto; pedido com vários produtos = "lote", marcado em
  `notes` com o prefixo `LOTE:<8 hex>:<N> itens | `;
- `order_number` = `id`; `total` = 0.0 (o app não controla preço por pedido);
- entrega hoje ou antes: o estoque é baixado na hora (`stock_reserved=1`); entrega futura:
  baixado pela tarefa de reserva quando a data chega;
- `status` é o PAGAMENTO; `situacao` é a ENTREGA (aberto | entregue | cancelado).
  Cancelar devolve o estoque reservado; entregue nunca devolve. Excluir só devolve o
  estoque de pedido ainda a entregar (aberto, com data de hoje em diante): excluir
  pedidos antigos para limpar a lista não "desfaz" bolos que já saíram.
"""

from __future__ import annotations

import os
import re
import uuid
from datetime import date, datetime

from api import ErroApi, Pedido, contexto, rota
from api.util import auditar, data_iso, ids, inteiro, like, linhas, qtd_texto, quantidade, texto, unidade_de

STATUS = ["Aguardando Pagamento", "Pagamento na Retirada", "Pago"]
SITUACOES = ["aberto", "entregue", "cancelado"]
_LOTE = re.compile(r"^LOTE:([0-9a-f]+):(\d+) itens \|\s?", re.I)


def _obs_segura(obs):
    """Observação digitada nunca pode começar com o marcador de lote (viraria um lote
    falso na lista e na exclusão)."""
    if obs and re.match(r"^\s*LOTE:", obs, re.I):
        return "Obs: " + obs.strip()
    return obs


def _validar_data_futura(d_iso: str) -> None:
    if date.fromisoformat(d_iso) < date.today():
        raise ErroApi("A data de entrega já passou. Escolha hoje ou uma data futura.", 400, "data_passada")


def separar_lote(notes):
    """("LOTE:ab12cd34:3 itens | sem açúcar") -> ("ab12cd34", 3, "sem açúcar")."""
    if not notes:
        return None, None, notes or ""
    m = _LOTE.match(notes)
    if not m:
        return None, None, notes
    return m.group(1), int(m.group(2)), notes[m.end():]


NOMES_DIAS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]


def dias_fechados(db) -> list[int]:
    """Dias da semana sem expediente (0 = segunda … 6 = domingo), da tabela company."""
    try:
        r = db.query("SELECT closed_weekdays FROM company WHERE id=1")
        bruto = (r[0]["closed_weekdays"] if r else "0") or ""
    except Exception:
        bruto = "0"
    return sorted({int(x) for x in str(bruto).split(",") if x.strip().isdigit() and 0 <= int(x) <= 6})


def _validar_entrega(d_iso: str) -> None:
    dia = date.fromisoformat(d_iso).weekday()
    if dia in dias_fechados(contexto.db()):
        raise ErroApi(f"{NOMES_DIAS[dia]} não há expediente, pedido negado.", 400, "dia_fechado")


# ---------------------------------------------------------------- listagem


@rota("GET", "/pedidos")
def listar_pedidos(p: Pedido) -> dict:
    """Filtro por data de entrega (de/ate opcionais) e texto (cliente, produto, status,
    etiqueta, observação). Itens de um lote aparecem numa linha só, com todos os ids."""
    db = contexto.db()
    where, params = ["1=1"], []
    if p.q("de"):
        where.append("DATE(o.delivery_date) >= ?")
        params.append(data_iso(p.q("de"), "de"))
    if p.q("ate"):
        where.append("DATE(o.delivery_date) <= ?")
        params.append(data_iso(p.q("ate"), "ate"))
    termo = p.q("q").strip()
    if termo:
        where.append("(" + " OR ".join(f"{c} LIKE ? ESCAPE '\\'" for c in
                     ("c.name", "p.name", "o.status", "COALESCE(o.label,'Comum')", "COALESCE(o.notes,'')")) + ")")
        params += [like(termo)] * 5
    situacao = p.q("situacao")
    if situacao in SITUACOES:
        where.append("COALESCE(o.situacao,'aberto') = ?")
        params.append(situacao)
    elif situacao != "todas":
        where.append("COALESCE(o.situacao,'aberto') != 'cancelado'")  # padrão: esconde cancelados
    rows = db.query(f"""
        SELECT
            MIN(o.id) AS id,
            GROUP_CONCAT(o.id) AS ids,
            COUNT(*) AS itens,
            o.customer_id, c.name AS cliente, c.phone AS telefone,
            CASE WHEN o.notes LIKE 'LOTE:%' THEN GROUP_CONCAT(p.name || ' ' || CASE WHEN p.unit = 'kg' THEN REPLACE(CAST(o.quantity AS TEXT), '.', ',') || ' kg' ELSE '×' || CAST(o.quantity AS INTEGER) END, ', ') ELSE p.name END AS produto,
            CASE WHEN o.notes LIKE 'LOTE:%' THEN NULL ELSE p.size END AS produto_size,
            ROUND(SUM(o.quantity), 3) AS quantity,
            -- unidade da linha: "kg" se todos os itens são em kg, "un" se todos em unidade, "misto" no lote com os dois
            CASE WHEN MIN(COALESCE(p.unit,'un')) = MAX(COALESCE(p.unit,'un')) THEN MIN(COALESCE(p.unit,'un')) ELSE 'misto' END AS unit,
            o.delivery_date, o.status,
            COALESCE(o.situacao, 'aberto') AS situacao,
            COALESCE(o.label, 'Comum') AS label,
            l.color AS label_color,
            o.notes,
            MIN(o.created_at) AS created_at,
            MIN(o.stock_reserved) AS stock_reserved
        FROM orders o
        JOIN customers c ON c.id = o.customer_id
        JOIN products p ON p.id = o.product_id
        LEFT JOIN labels l ON LOWER(l.name) = LOWER(COALESCE(o.label, 'Comum'))
        WHERE {" AND ".join(where)}
        GROUP BY o.customer_id, o.delivery_date, o.status, COALESCE(o.situacao, 'aberto'),
                 CASE WHEN o.notes LIKE 'LOTE:%' THEN SUBSTR(o.notes, 1, INSTR(o.notes, '|')) ELSE o.id END
        ORDER BY o.delivery_date DESC, MIN(o.created_at) DESC
    """, tuple(params))
    pedidos = []
    for r in rows:
        d = dict(r)
        _, _, d["obs"] = separar_lote(d["notes"])
        d["ids"] = [int(x) for x in str(d["ids"]).split(",")]
        d["lote"] = d["itens"] > 1 or bool(_LOTE.match(d["notes"] or ""))
        pedidos.append(d)
    return {"pedidos": pedidos, "status": STATUS, "situacoes": SITUACOES}


@rota("GET", "/pedidos/{id}")
def obter_pedido(p: Pedido) -> dict:
    """Um pedido e, se for de lote, todos os itens do lote."""
    db = contexto.db()
    r = db.query("SELECT o.*, c.name AS cliente FROM orders o JOIN customers c ON c.id=o.customer_id WHERE o.id=?", (p.id(),))
    if not r:
        raise ErroApi("Pedido não encontrado.", 404)
    o = dict(r[0])
    lote, _, o["obs"] = separar_lote(o["notes"])
    if lote:
        itens = db.query("""
            SELECT o.id, o.product_id, p.name AS produto, p.size AS produto_size, p.unit, o.quantity, o.stock_reserved
            FROM orders o JOIN products p ON p.id=o.product_id
            WHERE o.customer_id=? AND o.notes LIKE ? ORDER BY o.id
        """, (o["customer_id"], f"LOTE:{lote}:%"))
    else:
        itens = db.query("""
            SELECT o.id, o.product_id, p.name AS produto, p.size AS produto_size, p.unit, o.quantity, o.stock_reserved
            FROM orders o JOIN products p ON p.id=o.product_id WHERE o.id=?
        """, (o["id"],))
    o["itens"] = linhas(itens)
    return {"pedido": o}


# ---------------------------------------------------------------- criação


@rota("POST", "/pedidos")
def criar_pedido(p: Pedido) -> dict:
    """Cria um pedido com 1+ produtos (todos numa transação)."""
    db = contexto.db()
    cliente = inteiro(p.c("customer_id"), "Cliente", 1)
    entrega = data_iso(p.c("delivery_date"), "Entrega")
    _validar_data_futura(entrega)
    _validar_entrega(entrega)
    status = str(p.c("status") or STATUS[0])
    if status not in STATUS:
        raise ErroApi("Status inválido.")
    etiqueta = texto(p.c("label")) or "Comum"
    obs = _obs_segura(texto(p.c("notes")))
    produtos = p.c("products") or []
    if not isinstance(produtos, list) or not produtos:
        raise ErroApi("Adicione pelo menos um produto ao pedido!")
    itens = []
    for it in produtos:
        if not isinstance(it, dict):
            continue
        pid_ = inteiro(it.get("product_id"), "Produto", 1)
        # Produto vendido em kg aceita 1,5 (quilos); em unidade, só inteiro.
        itens.append((pid_, quantidade(it.get("quantity"), "Quantidade", unidade_de(db, pid_))))
    if not itens:
        raise ErroApi("Adicione pelo menos um produto ao pedido!")
    if not db.query("SELECT 1 FROM customers WHERE id=?", (cliente,)):
        raise ErroApi("Cliente não encontrado.")

    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    reservar = date.fromisoformat(entrega) <= date.today()
    if len(itens) > 1:
        notas = f"LOTE:{uuid.uuid4().hex[:8]}:{len(itens)} itens | {obs or ''}".strip()
    else:
        notas = obs
    criados = []
    with db.transaction():
        for pid, qtd in itens:
            if not db.query("SELECT 1 FROM products WHERE id=?", (pid,)):
                raise ErroApi("Produto não encontrado!")
            cur = db.execute(
                "INSERT INTO orders(order_number, customer_id, product_id, quantity, delivery_date, total, status, label, notes, stock_reserved, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (None, cliente, pid, qtd, entrega, 0.0, status, etiqueta, notas, 1 if reservar else 0, agora),
            )
            oid = cur.lastrowid
            db.execute("UPDATE orders SET order_number=? WHERE id=?", (oid, oid))
            if reservar:
                db.execute("UPDATE products SET stock = ROUND(stock - ?, 3) WHERE id=?", (qtd, pid))
                db.execute(
                    "INSERT INTO stock_movements(product_id, type, quantity, reason, order_id, created_at, created_by) VALUES (?,?,?,?,?,?,?)",
                    (pid, "saida", qtd, "Pedido", oid, agora, db.current_user or "unknown"),
                )
            auditar(db, "order", oid, "create", f"cust={cliente},prod={pid},qty={qtd},total=0.0,reserved={reservar}")
            criados.append(oid)
    return {"ids": criados}


# ---------------------------------------------------------------- edição


@rota("PUT", "/pedidos/{id}")
def atualizar_pedido(p: Pedido) -> dict:
    """Edita um item. Mantém as regras de estoque do app: ajusta pela diferença se já
    reservado, devolve/baixa ao trocar de produto, reserva se a data passa a ser hoje ou
    antes. Mover um pedido reservado para o futuro NÃO devolve o estoque.

    `situacao`: cancelar devolve o estoque reservado (e a reserva automática passa a
    ignorar o pedido); reabrir reserva de novo se a data já chegou; marcar entregue
    reserva sempre (o produto saiu)."""
    db = contexto.db()
    oid = p.id()
    r = db.query("SELECT * FROM orders WHERE id=?", (oid,))
    if not r:
        raise ErroApi("Pedido não encontrado.", 404)
    a = r[0]
    cliente = inteiro(p.c("customer_id", a["customer_id"]), "Cliente", 1)
    pid = inteiro(p.c("product_id", a["product_id"]), "Produto", 1)
    if pid != a["product_id"] and not db.query("SELECT 1 FROM products WHERE id=?", (pid,)):
        raise ErroApi("Produto não encontrado!")
    unidade = unidade_de(db, pid)
    qtd = quantidade(p.c("quantity", a["quantity"]), "Quantidade", unidade)
    entrega = data_iso(p.c("delivery_date", a["delivery_date"]), "Entrega")
    # Só confere o dia quando a data muda: um pedido antigo num dia que virou "sem
    # expediente" ainda pode ter o pagamento, a etiqueta etc. alterados.
    if entrega != data_iso(a["delivery_date"], "Entrega"):
        _validar_data_futura(entrega)
        _validar_entrega(entrega)
    status = str(p.c("status", a["status"]))
    if status not in STATUS and status != a["status"]:
        raise ErroApi("Status inválido.")
    etiqueta = texto(p.c("label", a["label"])) or "Comum"
    sit_antes = (a["situacao"] if "situacao" in a.keys() else None) or "aberto"
    situacao = str(p.c("situacao", sit_antes))
    if situacao not in SITUACOES:
        raise ErroApi("Situação inválida.")
    # Observação editada sem o marcador de lote; o marcador original é preservado.
    lote, n, _ = separar_lote(a["notes"])
    obs = _obs_segura(texto(p.c("notes"))) if "notes" in p.corpo else separar_lote(a["notes"])[2] or None
    notas = f"LOTE:{lote}:{n} itens | {obs or ''}".strip() if lote else obs

    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    reservado = int(a["stock_reserved"] or 0) == 1
    cancelar = situacao == "cancelado"
    reservar_agora = (not reservado and not cancelar
                      and (date.fromisoformat(entrega) <= date.today() or situacao == "entregue"))
    usuario = db.current_user or "unknown"

    def mov(prod, tipo, qtd_, motivo):
        db.execute("INSERT INTO stock_movements(product_id, type, quantity, reason, order_id, created_at, created_by) VALUES (?,?,?,?,?,?,?)",
                   (prod, tipo, qtd_, motivo, oid, agora, usuario))

    with db.transaction():
        if reservado and cancelar:
            db.execute("UPDATE products SET stock = ROUND(stock + ?, 3) WHERE id=?", (a["quantity"], a["product_id"]))
            mov(a["product_id"], "entrada", a["quantity"], "Pedido cancelado")
            reservado = False
        elif reservado and pid == a["product_id"]:
            delta = round(qtd - a["quantity"], 3)
            if delta:
                db.execute("UPDATE products SET stock = ROUND(stock - ?, 3) WHERE id=?", (delta, pid))
                mov(pid, "saida" if delta > 0 else "entrada", abs(delta), "Ajuste pedido")
        elif reservado:
            db.execute("UPDATE products SET stock = ROUND(stock + ?, 3) WHERE id=?", (a["quantity"], a["product_id"]))
            mov(a["product_id"], "entrada", a["quantity"], "Ajuste pedido")
            db.execute("UPDATE products SET stock = ROUND(stock - ?, 3) WHERE id=?", (qtd, pid))
            mov(pid, "saida", qtd, "Ajuste pedido")
        elif reservar_agora:
            db.execute("UPDATE products SET stock = ROUND(stock - ?, 3) WHERE id=?", (qtd, pid))
            mov(pid, "saida", qtd, "Pedido entregue" if situacao == "entregue" else "Pedido (reservado ao editar)")
        db.execute(
            "UPDATE orders SET customer_id=?, product_id=?, quantity=?, delivery_date=?, total=?, status=?, label=?, notes=?, stock_reserved=?, situacao=? WHERE id=?",
            (cliente, pid, qtd, entrega, a["total"] or 0.0, status, etiqueta, notas, 1 if (reservado or reservar_agora) else 0, situacao, oid),
        )
        mudancas = []
        if cliente != a["customer_id"]:
            nomes = {x["id"]: x["name"] for x in db.query("SELECT id, name FROM customers WHERE id IN (?,?)", (a["customer_id"], cliente))}
            mudancas.append(f"customer: '{nomes.get(a['customer_id'], '')}' -> '{nomes.get(cliente, '')}'")
        if pid != a["product_id"]:
            nomes = {x["id"]: x["name"] for x in db.query("SELECT id, name FROM products WHERE id IN (?,?)", (a["product_id"], pid))}
            mudancas.append(f"product: '{nomes.get(a['product_id'], '')}' -> '{nomes.get(pid, '')}'")
        if qtd != a["quantity"]:
            mudancas.append(f"quantity: {qtd_texto(a['quantity'], unidade)} -> {qtd_texto(qtd, unidade)}")
        if entrega != a["delivery_date"]:
            mudancas.append(f"delivery_date: '{a['delivery_date']}' -> '{entrega}'")
        if status != a["status"]:
            mudancas.append(f"status: '{a['status']}' -> '{status}'")
        if situacao != sit_antes:
            mudancas.append(f"situacao: '{sit_antes}' -> '{situacao}'")
        if etiqueta != (a["label"] or "Comum"):
            mudancas.append(f"label: '{a['label'] or 'Comum'}' -> '{etiqueta}'")
        if (notas or None) != (a["notes"] or None):
            mudancas.append(f"notes: '{a['notes'] or ''}' -> '{notas or ''}'")
        auditar(db, "order", oid, "update", "; ".join(mudancas) or None)
    return {"ok": True}


@rota("PUT", "/pedidos/lote")
def atualizar_lote(p: Pedido) -> dict:
    """Muda pagamento, situação, etiqueta, entrega ou observação de todos os itens das
    linhas marcadas (itens de lote incluídos)."""
    db = contexto.db()
    alvo = expandir_lotes(db, ids(p.c("ids")))
    campos = {k: p.c(k) for k in ("status", "situacao", "label", "delivery_date", "notes") if k in p.corpo}
    if not campos:
        raise ErroApi("Nada para alterar.")
    with db.transaction():  # tudo ou nada (a transação é reentrante)
        for oid in alvo:
            atualizar_pedido(Pedido("PUT", "", {}, campos, {"id": str(oid)}, p.usuario))
    return {"ok": True, "itens": len(alvo)}


@rota("POST", "/pedidos/editar")
def editar_pedido_completo(p: Pedido) -> dict:
    """Salva o diálogo de edição de uma linha (1 item ou lote) numa transação:
    campos comuns (cliente, entrega, status, etiqueta, observação) + produto/quantidade
    de cada item."""
    db = contexto.db()
    comuns = {k: p.c(k) for k in ("customer_id", "delivery_date", "status", "label", "notes") if k in p.corpo}
    itens = p.c("itens") or []
    if not isinstance(itens, list) or not itens:
        raise ErroApi("Nenhum item para salvar.")
    with db.transaction():
        for it in itens:
            if not isinstance(it, dict):
                raise ErroApi("Item inválido.")
            corpo = {**comuns, "product_id": it.get("product_id"), "quantity": it.get("quantity")}
            atualizar_pedido(Pedido("PUT", "", {}, corpo, {"id": str(inteiro(it.get("id"), "id"))}, p.usuario))
    return {"ok": True}


# ---------------------------------------------------------------- exclusão


def expandir_lotes(db, alvo: list[int]) -> list[int]:
    """Inclui todos os itens do lote de cada pedido selecionado."""
    todos: list[int] = []
    for oid in alvo:
        r = db.query("SELECT notes, customer_id FROM orders WHERE id=?", (oid,))
        if not r:
            continue
        lote, _, _ = separar_lote(r[0]["notes"])
        if lote:
            irmaos = [x["id"] for x in db.query("SELECT id FROM orders WHERE customer_id=? AND notes LIKE ? ORDER BY id",
                                                (r[0]["customer_id"], f"LOTE:{lote}:%"))]
        else:
            irmaos = [oid]
        todos += [i for i in irmaos if i not in todos]
    return todos


def _devolve_ao_excluir(o) -> bool:
    """Excluir devolve o estoque só de pedido reservado que ainda vai ser entregue
    (aberto, entrega de hoje em diante). Entregue, cancelado (já devolveu) ou com a
    data passada (presume-se entregue) não devolvem."""
    if o["product_id"] is None or int(o["stock_reserved"] or 0) != 1:
        return False
    if (o["situacao"] or "aberto") != "aberto":
        return False
    return data_iso(o["delivery_date"], "Entrega") >= date.today().isoformat()


_SQL_EXCLUSAO = "SELECT product_id, quantity, stock_reserved, delivery_date, situacao FROM orders WHERE id=?"


@rota("POST", "/pedidos/excluir")
def excluir_pedidos(p: Pedido) -> dict:
    """Apaga os pedidos. Pedido a entregar devolve o estoque e as movimentações dele
    somem (o efeito se anula). Pedido entregue ou antigo não devolve: as saídas ficam
    no histórico de movimentações, só desvinculadas do pedido apagado."""
    db = contexto.db()
    alvo = expandir_lotes(db, ids(p.c("ids")))
    devolvidos = 0
    with db.transaction():
        for oid in alvo:
            r = db.query(_SQL_EXCLUSAO, (oid,))
            if not r:
                continue
            if _devolve_ao_excluir(r[0]):
                db.execute("UPDATE products SET stock = ROUND(stock + ?, 3) WHERE id=?", (r[0]["quantity"], r[0]["product_id"]))
                db.execute("DELETE FROM stock_movements WHERE order_id=?", (oid,))
                devolvidos += r[0]["quantity"]
            else:
                db.execute("UPDATE stock_movements SET order_id=NULL, reason=COALESCE(reason,'Pedido') || ' #' || ? || ' (excluído)' "
                           "WHERE order_id=?", (oid, oid))
            db.execute("DELETE FROM orders WHERE id=?", (oid,))
            auditar(db, "order", oid, "delete")
    return {"excluidos": len(alvo), "devolvidos": devolvidos}


@rota("POST", "/pedidos/excluir/contar", escreve=False)
def contar_exclusao(p: Pedido) -> dict:
    """Para a confirmação: quantos itens serão excluídos e quantas unidades voltam."""
    db = contexto.db()
    alvo = expandir_lotes(db, ids(p.c("ids")))
    devolve = 0
    for oid in alvo:
        r = db.query(_SQL_EXCLUSAO, (oid,))
        if r and _devolve_ao_excluir(r[0]):
            devolve += r[0]["quantity"]
    return {"itens": len(alvo), "devolve": devolve}


# ---------------------------------------------------------------- relatório


def _logo_empresa(db) -> tuple[str, str | None]:
    from core.paths import RESOURCE_DIR
    r = db.query("SELECT name, logo_path FROM company WHERE id=1")
    nome = (r[0]["name"] if r else None) or "Confeitaria"
    logo = r[0]["logo_path"] if r else None
    if not logo or not os.path.isfile(logo):
        for c in (os.path.join(RESOURCE_DIR, "assets", "icons", "logo.png"),):
            if os.path.isfile(c):
                logo = c
                break
        else:
            logo = None
    return nome, logo


@rota("GET", "/pedidos/relatorio/opcoes")
def opcoes_relatorio(p: Pedido) -> dict:
    db = contexto.db()
    # Por cadastro (id), não por nome: clientes homônimos não se misturam no relatório.
    clientes = [{"id": r["id"], "nome": r["name"], "telefone": r["phone"]} for r in db.query(
        "SELECT DISTINCT c.id, c.name, c.phone FROM orders o JOIN customers c ON c.id = o.customer_id "
        "WHERE c.name IS NOT NULL AND c.name != '' ORDER BY c.name, c.phone")]
    etiquetas = [r["name"] for r in db.query("SELECT DISTINCT name FROM labels WHERE name IS NOT NULL AND name != '' ORDER BY name")]
    return {"clientes": clientes, "etiquetas": etiquetas}


@rota("POST", "/pedidos/relatorio", escreve=False)
def relatorio(p: Pedido) -> dict:
    """PDF (reportlab, A4 paisagem) ou Excel (openpyxl) dos pedidos do período."""
    from core.formatting import format_date, format_size

    db = contexto.db()
    de, ate = data_iso(p.c("de"), "de"), data_iso(p.c("ate"), "ate")
    cliente_id = p.c("cliente_id")
    cliente_id = inteiro(cliente_id, "Cliente", 1) if cliente_id not in (None, "") else None
    cliente = None
    if cliente_id:
        r = db.query("SELECT name FROM customers WHERE id=?", (cliente_id,))
        cliente = r[0]["name"] if r else None
    etiqueta = texto(p.c("etiqueta"))
    formato = p.c("formato")
    caminho = str(p.c("caminho") or "")
    if formato not in ("pdf", "xlsx") or not caminho.lower().endswith("." + formato):
        raise ErroApi("Formato ou arquivo inválido.")
    sql = """
        SELECT c.name AS cliente, p.name AS produto, p.size AS produto_size, p.unit, o.quantity, o.delivery_date, o.status,
               COALESCE(o.label, 'Comum') AS label, o.notes
        FROM orders o JOIN customers c ON c.id = o.customer_id JOIN products p ON p.id = o.product_id
        WHERE DATE(o.delivery_date) BETWEEN ? AND ? AND COALESCE(o.situacao, 'aberto') != 'cancelado'"""
    params: list = [de, ate]
    if cliente_id:
        sql += " AND o.customer_id = ?"
        params.append(cliente_id)
    if etiqueta == "Comum":
        sql += " AND (o.label IS NULL OR o.label = 'Comum')"
    elif etiqueta:
        sql += " AND o.label = ?"
        params.append(etiqueta)
    rows = db.query(sql + " ORDER BY o.delivery_date ASC, o.created_at ASC", tuple(params))
    if not rows:
        raise ErroApi("Nenhum pedido encontrado com os filtros selecionados.", 404, "vazio")

    empresa, logo = _logo_empresa(db)
    sub = f"Período: {format_date(de)} a {format_date(ate)}"
    if cliente:
        sub += f" | Cliente: {cliente}"
    if etiqueta:
        sub += f" | Etiqueta: {etiqueta}"
    cab = ["#", "Cliente", "Produto", "Tamanho", "Qtd", "Entrega", "Status", "Etiqueta", "Observações"]
    dados = [[i, r["cliente"], r["produto"], format_size(r["produto_size"]), qtd_texto(r["quantity"], r["unit"]),
              format_date(r["delivery_date"]), r["status"], r["label"] or "Comum", separar_lote(r["notes"])[2]]
             for i, r in enumerate(rows, 1)]
    # Total separado por unidade: somar unidades com quilos não faz sentido.
    un = sum(int(r["quantity"]) for r in rows if (r["unit"] or "un") != "kg")
    kg = round(sum(float(r["quantity"]) for r in rows if r["unit"] == "kg"), 3)
    total = " + ".join(x for x in (f"{un} un" if un or not kg else "", qtd_texto(kg, "kg") if kg else "") if x)
    gerado = f"Relatório gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}"
    if formato == "pdf":
        _relatorio_pdf(caminho, empresa, logo, sub, cab, dados, total, gerado)
    else:
        _relatorio_xlsx(caminho, empresa, logo, sub, cab, dados, total, gerado)
    return {"caminho": caminho, "linhas": len(dados)}


def _relatorio_pdf(caminho, empresa, logo, sub, cab, dados, total, gerado):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    doc = SimpleDocTemplate(caminho, pagesize=landscape(A4), topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm, title="Relatório de Pedidos")
    est = getSampleStyleSheet()
    titulo = ParagraphStyle("T", parent=est["Heading1"], fontSize=18, textColor=colors.HexColor("#1f2937"), alignment=TA_CENTER, spaceAfter=6)
    sub_e = ParagraphStyle("S", parent=est["Normal"], fontSize=11, textColor=colors.HexColor("#6b7280"), alignment=TA_CENTER)
    cel = ParagraphStyle("C", parent=est["Normal"], fontSize=9, leading=11, textColor=colors.HexColor("#374151"))
    rod = ParagraphStyle("R", parent=est["Normal"], fontSize=8, textColor=colors.HexColor("#9ca3af"), alignment=TA_CENTER)
    story = []
    if logo:
        try:
            from reportlab.lib.utils import ImageReader
            w, h = ImageReader(logo).getSize()
            esc = min(3 * cm / w, 3 * cm / h)
            img = Image(logo, width=w * esc, height=h * esc)
            img.hAlign = "LEFT"
            story += [img, Spacer(1, 0.3 * cm)]
        except Exception:
            pass
    story += [Paragraph(escape(empresa), titulo), Paragraph("Relatório de Pedidos", sub_e), Paragraph(escape(sub), sub_e), Spacer(1, 0.5 * cm)]
    tabela = [cab] + [[Paragraph(escape(str(v)), cel) for v in linha] for linha in dados]
    tabela.append(["", "", "", "TOTAL:", str(total), "", "", "", ""])
    t = Table(tabela, colWidths=[w * cm for w in [1.2, 3.5, 3.5, 2, 1.2, 2.2, 3, 2.5, 4]], repeatRows=1)
    n = len(tabela)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, n - 2), 0.5, colors.HexColor("#e5e7eb")),
        ("FONTNAME", (0, n - 1), (-1, n - 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, n - 1), (-1, n - 1), 10),
        ("BACKGROUND", (0, n - 1), (-1, n - 1), colors.HexColor("#f3f4f6")),
        ("BOX", (0, n - 1), (-1, n - 1), 1, colors.HexColor("#1f2937")),
    ]
    for i in range(1, n - 1):
        estilo.append(("BACKGROUND", (0, i), (-1, i), colors.white if i % 2 else colors.HexColor("#f9fafb")))
    t.setStyle(TableStyle(estilo))
    story += [t, Spacer(1, 1 * cm), Paragraph(gerado, rod)]
    doc.build(story)


def _relatorio_xlsx(caminho, empresa, logo, sub, cab, dados, total, gerado):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório de Pedidos"
    primeira = 1
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(logo)
            esc = 100 / max(img.width, img.height)
            img.width, img.height = img.width * esc, img.height * esc
            ws.add_image(img, "A1")
            primeira = 6
        except Exception:
            pass
    for i, (txt, fonte) in enumerate([(empresa, Font(size=18, bold=True, color="1f2937")),
                                      ("Relatório de Pedidos", Font(size=13, bold=True, color="374151")),
                                      (sub, Font(size=11, color="6b7280"))]):
        lin = primeira + i
        ws.merge_cells(start_row=lin, start_column=1, end_row=lin, end_column=9)
        c = ws.cell(lin, 1, txt)
        c.font, c.alignment = fonte, Alignment(horizontal="center")
    fina = Side(style="thin", color="e5e7eb")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    lin_cab = primeira + 4
    for j, nome in enumerate(cab, 1):
        c = ws.cell(lin_cab, j, nome)
        c.fill = PatternFill("solid", fgColor="1f2937")
        c.font = Font(color="FFFFFF", bold=True, size=11)
        c.alignment, c.border = Alignment(horizontal="center", vertical="center"), borda
    for i, linha in enumerate(dados):
        lin = lin_cab + 1 + i
        for j, v in enumerate(linha, 1):
            c = ws.cell(lin, j, v)
            c.border = borda
            c.fill = PatternFill("solid", fgColor="FFFFFF" if i % 2 == 0 else "f9fafb")
            c.alignment = Alignment(horizontal="center", vertical="center") if j in (1, 5, 6) else Alignment(vertical="top", wrap_text=True)
    lin_tot = lin_cab + 1 + len(dados)
    media = Side(style="medium", color="1f2937")
    for j in range(1, 10):
        c = ws.cell(lin_tot, j)
        c.fill = PatternFill("solid", fgColor="f3f4f6")
        c.border = Border(left=media, right=media, top=media, bottom=media)
        c.font = Font(bold=True, size=11)
    ws.cell(lin_tot, 4, "TOTAL:")
    ws.cell(lin_tot, 5, total)
    for letra, w in zip("ABCDEFGHI", [8, 25, 25, 12, 8, 15, 20, 15, 35]):
        ws.column_dimensions[letra].width = w
    c = ws.cell(lin_tot + 2, 1, gerado)
    c.font = Font(size=9, italic=True, color="9ca3af")
    wb.save(caminho)
