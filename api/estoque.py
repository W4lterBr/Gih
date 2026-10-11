# -*- coding: utf-8 -*-
"""Produtos (catálogo e pronta entrega) e movimentações de estoque."""

from __future__ import annotations

from datetime import datetime

from api import ErroApi, Pedido, contexto, rota
from api.util import agora, auditar, data_iso, diff, ids, inteiro, linhas, texto
from core.stock_categories import MOVEMENT_CATEGORIES, MOVEMENT_CATEGORY_COLORS

# ---------------------------------------------------------------- produtos


@rota("GET", "/produtos")
def listar_produtos(p: Pedido) -> dict:
    """`ordem=pronta` coloca esgotados no fim (tela Pronta entrega)."""
    db = contexto.db()
    termo = p.q("q").strip()
    ordem = "(stock <= 0), name" if p.q("ordem") == "pronta" else "name"
    if termo:
        like = f"%{termo}%"
        rows = db.query(f"SELECT * FROM products WHERE name LIKE ? OR description LIKE ? ORDER BY {ordem}", (like, like))
    else:
        rows = db.query(f"SELECT * FROM products ORDER BY {ordem}")
    return {"produtos": linhas(rows)}


@rota("GET", "/produtos/estoque-baixo")
def estoque_baixo(p: Pedido) -> dict:
    rows = contexto.db().query(
        "SELECT id, name, stock, min_stock FROM products WHERE stock <= min_stock AND min_stock > 0 ORDER BY name"
    )
    return {"produtos": linhas(rows)}


def _campos_produto(p: Pedido) -> dict:
    nome = texto(p.c("name"))
    if not nome:
        raise ErroApi("Nome é obrigatório.")
    tamanhos = p.c("sizes") or []
    if not isinstance(tamanhos, list):
        raise ErroApi("Tamanhos inválidos.")
    tamanhos = [str(t).strip() for t in tamanhos if str(t).strip()]
    try:
        preco = round(float(p.c("price") or 0), 2)
    except (TypeError, ValueError):
        raise ErroApi("Preço inválido.")
    if preco < 0 or preco > 999999.99:
        raise ErroApi("Preço inválido.")
    return {
        "name": nome,
        "description": texto(p.c("description")),
        "sizes": tamanhos,
        "stock": inteiro(p.c("stock", 0), "Estoque", 0, 1_000_000),
        "min_stock": inteiro(p.c("min_stock", 0), "Estoque mínimo", 0, 100_000),
        "price": preco,
    }


def _nome_variante(nome: str, tamanho: str) -> str:
    # "15" -> "Bolo 15cm"; quem já digitou "15 cm"/"15cm" não ganha "cm" duplicado.
    t = tamanho.strip()
    return f"{nome} {t}" if t.lower().endswith("cm") else f"{nome} {t}cm"


def _inserir_produto(db, nome: str, f: dict, tamanho, detalhe_preco: bool) -> int:
    cur = db.execute(
        "INSERT INTO products(name, description, size, stock, min_stock, price) VALUES (?,?,?,?,?,?)",
        (nome, f["description"], tamanho, f["stock"], f["min_stock"], f["price"]),
    )
    novo = cur.lastrowid
    if f["stock"] > 0:
        db.execute(
            "INSERT INTO stock_movements(product_id, type, category, quantity, reason, order_id, created_at, created_by) VALUES (?,?,?,?,?,?,?,?)",
            (novo, "entrada", "Ajuste de Inventário", f["stock"], "Cadastro inicial", None, agora(), db.current_user or "unknown"),
        )
    det = f"name={nome},stock={f['stock']}" + (f",price={f['price']}" if detalhe_preco else "")
    auditar(db, "product", novo, "create", det)
    return novo


@rota("POST", "/produtos")
def criar_produto(p: Pedido) -> dict:
    """Mais de um tamanho: cria um produto por tamanho ("Bolo 15cm", "Bolo 20cm")."""
    db = contexto.db()
    f = _campos_produto(p)
    criados = []
    with db.transaction():
        if len(f["sizes"]) > 1:
            for t in f["sizes"]:
                criados.append(_inserir_produto(db, _nome_variante(f["name"], t), f, t, True))
        else:
            criados.append(_inserir_produto(db, f["name"], f, f["sizes"][0] if f["sizes"] else None, False))
    return {"ids": criados}


@rota("PUT", "/produtos/{id}")
def atualizar_produto(p: Pedido) -> dict:
    """Edita; com vários tamanhos o atual fica com o 1º e os demais viram produtos novos.
    Mudança de estoque gera movimentação 'Ajuste de Inventário'."""
    db = contexto.db()
    pid = p.id()
    r = db.query("SELECT * FROM products WHERE id=?", (pid,))
    if not r:
        raise ErroApi("Produto não encontrado.", 404)
    a = r[0]
    f = _campos_produto(p)
    novos = []
    with db.transaction():
        if len(f["sizes"]) > 1:
            nome_atual, tamanho_atual = _nome_variante(f["name"], f["sizes"][0]), f["sizes"][0]
        else:
            nome_atual, tamanho_atual = f["name"], (f["sizes"][0] if f["sizes"] else None)
        db.execute(
            "UPDATE products SET name=?, description=?, size=?, stock=?, min_stock=?, price=? WHERE id=?",
            (nome_atual, f["description"], tamanho_atual, f["stock"], f["min_stock"], f["price"], pid),
        )
        delta = f["stock"] - int(a["stock"] or 0)
        if delta:
            db.execute(
                "INSERT INTO stock_movements(product_id, type, category, quantity, reason, order_id, created_at, created_by) VALUES (?,?,?,?,?,?,?,?)",
                (pid, "entrada" if delta > 0 else "saida", "Ajuste de Inventário", abs(delta), "Edição de produto", None, agora(), db.current_user or "unknown"),
            )
        for t in f["sizes"][1:]:
            novos.append(_inserir_produto(db, _nome_variante(f["name"], t), f, t, True))
        partes = [
            diff([("name", a["name"], nome_atual), ("description", a["description"], f["description"])]),
            diff([("stock", a["stock"], f["stock"]), ("min_stock", a["min_stock"], f["min_stock"])], aspas=False),
        ]
        auditar(db, "product", pid, "update", "; ".join(x for x in partes if x) or None)
    return {"ok": True, "novos": novos}


@rota("POST", "/produtos/excluir")
def excluir_produtos(p: Pedido) -> dict:
    db = contexto.db()
    excluidos, bloqueados = 0, []
    with db.transaction():
        for pid in ids(p.c("ids")):
            if db.query("SELECT COUNT(*) AS c FROM orders WHERE product_id=?", (pid,))[0]["c"] > 0:
                r = db.query("SELECT name FROM products WHERE id=?", (pid,))
                bloqueados.append(r[0]["name"] if r else f"#{pid}")
                continue
            db.execute("DELETE FROM stock_movements WHERE product_id=?", (pid,))
            db.execute("DELETE FROM production_items WHERE product_id=?", (pid,))
            db.execute("DELETE FROM products WHERE id=?", (pid,))
            auditar(db, "product", pid, "delete")
            excluidos += 1
    return {"excluidos": excluidos, "bloqueados": bloqueados}


# ---------------------------------------------------------------- movimentações

_PERIODOS = {
    "hoje": "DATE(sm.created_at) = DATE('now','localtime')",
    "7d": "DATE(sm.created_at) >= DATE('now','localtime','-7 days')",
    "30d": "DATE(sm.created_at) >= DATE('now','localtime','-30 days')",
    "tudo": "1=1",
}
_FILTROS = {
    "entrada": "sm.type = 'entrada'",
    "saida": "sm.type = 'saida'",
    "producao": "(sm.category = 'Produção' OR (sm.category IS NULL AND sm.reason LIKE 'Produção%'))",
    "perda": "sm.category = 'Perda/Quebra'",
    "todas": "1=1",
}


def _where_movimentos(p: Pedido) -> tuple[str, tuple]:
    partes = [_PERIODOS.get(p.q("periodo", "30d"), _PERIODOS["30d"]), _FILTROS.get(p.q("filtro", "todas"), "1=1")]
    params: tuple = ()
    termo = p.q("q").strip()
    if termo:
        partes.append("p.name LIKE ?")
        params = (f"%{termo}%",)
    return " AND ".join(partes), params


@rota("GET", "/movimentos")
def listar_movimentos(p: Pedido) -> dict:
    db = contexto.db()
    where, params = _where_movimentos(p)
    rows = db.query(f"""
        SELECT sm.id, sm.type, sm.category, sm.quantity, sm.reason,
               sm.created_at, sm.created_by, sm.order_id, sm.product_id, p.name AS product_name
        FROM stock_movements sm
        JOIN products p ON p.id = sm.product_id
        WHERE {where}
        ORDER BY sm.created_at DESC, sm.id DESC
        LIMIT 500
    """, params)
    resumo = db.query("""
        SELECT
            COALESCE(SUM(CASE WHEN type='entrada' AND DATE(created_at)=DATE('now','localtime') THEN quantity ELSE 0 END), 0) AS entradas_hoje,
            COALESCE(SUM(CASE WHEN type='saida' AND DATE(created_at)=DATE('now','localtime') THEN quantity ELSE 0 END), 0) AS saidas_hoje,
            COALESCE(SUM(CASE WHEN (category='Produção' OR (category IS NULL AND reason LIKE 'Produção%'))
                            AND strftime('%Y-%m', created_at)=strftime('%Y-%m','now','localtime') THEN quantity ELSE 0 END), 0) AS producao_mes
        FROM stock_movements
    """)[0]
    perdas = db.query("""
        SELECT COALESCE(SUM(sm.quantity), 0) AS qtd, COALESCE(SUM(sm.quantity * COALESCE(p.price, 0)), 0) AS valor
        FROM stock_movements sm LEFT JOIN products p ON p.id = sm.product_id
        WHERE sm.category = 'Perda/Quebra' AND strftime('%Y-%m', sm.created_at) = strftime('%Y-%m','now','localtime')
    """)[0]
    return {
        "movimentos": linhas(rows),
        "resumo": {**dict(resumo), "perdas_qtd": perdas["qtd"], "perdas_valor": perdas["valor"]},
        "categorias": MOVEMENT_CATEGORIES,
        "cores": MOVEMENT_CATEGORY_COLORS,
    }


@rota("POST", "/movimentos")
def criar_movimento(p: Pedido) -> dict:
    db = contexto.db()
    tipo = p.c("type")
    if tipo not in ("entrada", "saida"):
        raise ErroApi("Tipo inválido.")
    categoria = str(p.c("category") or "")
    if categoria not in MOVEMENT_CATEGORIES["Entrada" if tipo == "entrada" else "Saída"]:
        raise ErroApi("Categoria inválida.")
    pid = inteiro(p.c("product_id"), "Produto", 1)
    qtd = inteiro(p.c("quantity"), "Quantidade", 1, 999_999)
    dia = data_iso(p.c("date") or datetime.now().date().isoformat(), "Data")
    if dia > datetime.now().date().isoformat():
        raise ErroApi("A data da movimentação não pode ser futura.")
    prod = db.query("SELECT name, stock FROM products WHERE id=?", (pid,))
    if not prod:
        raise ErroApi("Cadastre um produto antes de registrar uma movimentação.")
    quando = f"{dia} {datetime.now().strftime('%H:%M:%S')}"
    with db.transaction():
        if tipo == "entrada":
            db.execute("UPDATE products SET stock = stock + ? WHERE id=?", (qtd, pid))
        else:
            cur = db.execute("UPDATE products SET stock = stock - ? WHERE id=? AND stock >= ?", (qtd, pid, qtd))
            if cur.rowcount == 0:
                disp = db.query("SELECT stock FROM products WHERE id=?", (pid,))[0]["stock"]
                raise ErroApi(f"Estoque disponível: {disp} un. Não é possível dar saída de {qtd} un.", 409, "estoque_insuficiente")
        cur = db.execute(
            "INSERT INTO stock_movements(product_id, type, category, quantity, reason, order_id, created_at, created_by) VALUES (?,?,?,?,?,?,?,?)",
            (pid, tipo, categoria, qtd, texto(p.c("notes")), None, quando, db.current_user or "unknown"),
        )
        mid = cur.lastrowid
        auditar(db, "stock_movement", mid, "create", f"{tipo} | {categoria} | produto={prod[0]['name']} | qtd={qtd}")
    return {"id": mid, "produto": prod[0]["name"]}


@rota("POST", "/movimentos/excluir")
def excluir_movimentos(p: Pedido) -> dict:
    """Desfaz o efeito no estoque. Movimentos gerados por pedido não podem ser excluídos aqui."""
    db = contexto.db()
    excluidos, bloqueados = 0, 0
    with db.transaction():
        for mid in ids(p.c("ids")):
            r = db.query("SELECT product_id, type, quantity, order_id FROM stock_movements WHERE id=?", (mid,))
            if not r:
                continue
            m = r[0]
            if m["order_id"] is not None:
                bloqueados += 1
                continue
            delta = -m["quantity"] if m["type"] == "entrada" else m["quantity"]
            db.execute("UPDATE products SET stock = stock + ? WHERE id=?", (delta, m["product_id"]))
            db.execute("DELETE FROM stock_movements WHERE id=?", (mid,))
            auditar(db, "stock_movement", mid, "delete", f"revertido: {m['type']} {m['quantity']}un")
            excluidos += 1
    return {"excluidos": excluidos, "bloqueados": bloqueados}


@rota("POST", "/movimentos/exportar", escreve=False)
def exportar_movimentos(p: Pedido) -> dict:
    """Excel com os filtros atuais (sem limite de linhas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from core.formatting import format_datetime

    caminho = str(p.c("caminho") or "")
    if not caminho.lower().endswith(".xlsx"):
        raise ErroApi("Escolha um arquivo .xlsx.")
    filtro = Pedido("GET", "", {k: str(v) for k, v in (p.c("filtros") or {}).items()}, {})
    where, params = _where_movimentos(filtro)
    rows = contexto.db().query(f"""
        SELECT sm.created_at, p.name AS product_name, sm.type, sm.category, sm.quantity, sm.reason, sm.created_by
        FROM stock_movements sm JOIN products p ON p.id = sm.product_id
        WHERE {where} ORDER BY sm.created_at DESC
    """, params)
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimentações"
    cab = ["Data", "Produto", "Tipo", "Categoria", "Quantidade", "Motivo", "Usuário"]
    ws.append(cab)
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="8B5CF6")
        cel.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([format_datetime(r["created_at"]), r["product_name"], "Entrada" if r["type"] == "entrada" else "Saída",
                   r["category"] or "", r["quantity"], r["reason"] or "", r["created_by"] or ""])
    for col in ws.columns:
        largura = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(largura + 2, 12), 40)
    wb.save(caminho)
    return {"caminho": caminho, "linhas": len(rows)}
